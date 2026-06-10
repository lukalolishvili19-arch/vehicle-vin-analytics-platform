#!/usr/bin/env python3
"""Generate Excel workbook from fct_vehicle_inventory.csv (no pandas)."""

from __future__ import annotations

import csv
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT = Path(__file__).resolve().parents[1]
SRC = PROJECT / "powerbi" / "data" / "fct_vehicle_inventory.csv"
OUT = PROJECT / "excel" / "output" / "vehicle_analytics.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4788")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def load_rows() -> list[dict[str, str]]:
    with SRC.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def num(v, default=None):
    if v in ("", None):
        return default
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def write_sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18


def build_kpi(rows: list[dict]) -> list[list]:
    n = len(rows)
    grades = [num(r["grade"]) for r in rows if num(r.get("grade")) is not None]
    mileages = [int(num(r["mileage"], 0)) for r in rows]
    models = {r["model"] for r in rows}
    with_report = sum(1 for r in rows if r.get("has_condition_report") == "1")
    lights = Counter(r.get("primary_light", "UNKNOWN") for r in rows)
    return [
        ["Total Vehicles", n],
        ["Unique Models", len(models)],
        ["Avg Mileage", round(sum(mileages) / n) if n else 0],
        ["Avg Grade", round(sum(grades) / len(grades), 2) if grades else ""],
        ["Condition Report %", round(100 * with_report / n, 1) if n else 0],
        ["AS IS Count", sum(1 for r in rows if r.get("is_as_is") == "1")],
        ["Green Light Count", lights.get("GREEN", 0)],
        ["Red Light Count", lights.get("RED", 0)],
        ["Structural Damage", sum(1 for r in rows if r.get("has_structural_damage") == "1")],
    ]


def group_agg(rows: list[dict], key: str) -> list[list]:
    buckets: dict[str, list] = defaultdict(list)
    for r in rows:
        buckets[r.get(key) or "UNKNOWN"].append(r)
    out = []
    for k, items in sorted(buckets.items(), key=lambda x: -len(x[1])):
        grades = [num(r["grade"]) for r in items if num(r.get("grade")) is not None]
        mileages = [int(num(r["mileage"], 0)) for r in items]
        out.append([
            k,
            len(items),
            round(sum(mileages) / len(mileages)) if mileages else 0,
            round(sum(grades) / len(grades), 2) if grades else "",
        ])
    return out


def group_year(rows: list[dict]) -> list[list]:
    buckets: dict[int, list] = defaultdict(list)
    for r in rows:
        y = int(num(r.get("model_year"), 0))
        buckets[y].append(r)
    out = []
    for y in sorted(buckets.keys(), reverse=True):
        items = buckets[y]
        grades = [num(r["grade"]) for r in items if num(r.get("grade")) is not None]
        mileages = [int(num(r["mileage"], 0)) for r in items]
        out.append([
            y,
            len(items),
            round(sum(grades) / len(grades), 2) if grades else "",
            round(sum(mileages) / len(mileages)) if mileages else 0,
        ])
    return out


def detail_rows(rows: list[dict], limit: int = 555) -> tuple[list[str], list[list]]:
    cols = [
        "vin", "stock_number", "model_year", "model", "style", "exterior_color",
        "mileage", "grade", "primary_light", "mileage_bucket", "grade_tier",
        "has_condition_report", "is_as_is", "announcements",
    ]
    data = []
    for r in rows[:limit]:
        data.append([r.get(c, "") for c in cols])
    return cols, data


def main() -> None:
    if not SRC.exists():
        raise FileNotFoundError(f"Run ETL first: {SRC}")

    rows = load_rows()
    wb = Workbook()
    wb.remove(wb.active)

    sheets = [
        ("KPI Dashboard", ["Metric", "Value"], build_kpi(rows)),
        ("By Model", ["model", "vehicles", "avg_mileage", "avg_grade"], group_agg(rows, "model")),
        ("By Year", ["model_year", "count", "avg_grade", "avg_mileage"], group_year(rows)),
        ("By Color", ["exterior_color", "count"], [[k, v] for k, v in Counter(r.get("exterior_color", "") for r in rows).most_common()]),
        ("By Light", ["primary_light", "count"], [[k, v] for k, v in Counter(r.get("primary_light", "UNKNOWN") for r in rows).most_common()]),
        ("Mileage Bucket", ["mileage_bucket", "count"], [[k, v] for k, v in Counter(r.get("mileage_bucket", "") for r in rows).most_common()]),
    ]

    for name, headers, data in sheets:
        ws = wb.create_sheet(name)
        write_sheet(ws, headers, data)

    detail_headers, detail_data = detail_rows(rows)
    ws_detail = wb.create_sheet("Detail Data")
    write_sheet(ws_detail, detail_headers, detail_data)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Excel workbook created: {OUT}")
    print(f"  Rows: {len(rows)} | Sheets: {len(wb.sheetnames)}")


if __name__ == "__main__":
    main()
